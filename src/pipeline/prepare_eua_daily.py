from __future__ import annotations

import argparse
import io
import json
import zipfile
from datetime import date, datetime
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = ROOT / "data" / "raw" / "eua_daily.csv"
DEFAULT_METADATA = ROOT / "data" / "raw" / "eua_daily_metadata.json"


def _normalise_header(value: object) -> str:
    return " ".join(str(value or "").replace("�", " ").lower().split())


def _column_indices(header: tuple[object, ...]) -> dict[str, int]:
    labels = [_normalise_header(value) for value in header]

    def find(*needles: str) -> int:
        for index, label in enumerate(labels):
            if all(needle in label for needle in needles):
                return index
        raise RuntimeError(f"Could not find workbook column containing {needles!r}")

    return {
        "date": find("date"),
        "auction_name": find("auction name"),
        "contract": find("contract"),
        "status": find("status"),
        "price": find("auction price"),
        "volume": find("auction volume"),
    }


def _read_workbook(stream: object, source_name: str) -> list[dict[str, object]]:
    workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    if "Primary Market Auction" not in workbook.sheetnames:
        raise RuntimeError(f"{source_name} lacks the 'Primary Market Auction' sheet")
    sheet = workbook["Primary Market Auction"]
    header = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
    indices = _column_indices(header)
    records: list[dict[str, object]] = []
    for row in sheet.iter_rows(min_row=7, values_only=True):
        values = {name: row[index] for name, index in indices.items()}
        raw_date = values["date"]
        if isinstance(raw_date, datetime):
            raw_date = raw_date.date()
        if not isinstance(raw_date, date):
            continue
        price = pd.to_numeric(pd.Series([values["price"]]), errors="coerce").iloc[0]
        volume = pd.to_numeric(pd.Series([values["volume"]]), errors="coerce").iloc[0]
        records.append(
            {
                "date": pd.Timestamp(raw_date).normalize(),
                "auction_name": str(values["auction_name"] or ""),
                "contract": str(values["contract"] or ""),
                "status": str(values["status"] or ""),
                "price": float(price) if pd.notna(price) else np.nan,
                "volume": float(volume) if pd.notna(volume) else np.nan,
            }
        )
    workbook.close()
    return records


def _source_records(historical_zip: Path, current_xlsx: Path) -> tuple[list[dict[str, object]], list[str]]:
    records: list[dict[str, object]] = []
    source_names: list[str] = []
    with zipfile.ZipFile(historical_zip) as archive:
        for name in sorted(archive.namelist()):
            # The project starts in 2021; older archive members add no
            # information to the forecast and some are legacy .xls files.
            if not name.lower().endswith(".xlsx"):
                continue
            if not any(f"-{year}-data.xlsx" in name for year in range(2021, 2026)):
                continue
            source_names.append(name)
            records.extend(_read_workbook(io.BytesIO(archive.read(name)), name))
    source_names.append(current_xlsx.name)
    records.extend(_read_workbook(current_xlsx, current_xlsx.name))
    return records, source_names


def build(historical_zip: Path, current_xlsx: Path, output: Path, metadata_path: Path) -> pd.DataFrame:
    records, source_names = _source_records(historical_zip, current_xlsx)
    raw = pd.DataFrame(records)
    if raw.empty:
        raise RuntimeError("No auction records were found in the supplied files")

    source_rows = len(raw)
    filtered = raw[
        raw["contract"].eq("T3PA")
        & raw["status"].str.lower().eq("successful")
        & raw["price"].notna()
        & raw["volume"].gt(0)
    ].copy()
    if filtered.empty:
        raise RuntimeError("No successful T3PA allowance auctions were found")

    filtered["price_volume"] = filtered["price"] * filtered["volume"]
    daily = (
        filtered.groupby("date", as_index=False)
        .agg(price_volume=("price_volume", "sum"), volume=("volume", "sum"))
        .assign(eua_eur_tco2=lambda frame: frame["price_volume"] / frame["volume"])
        .loc[:, ["date", "eua_eur_tco2"]]
        .sort_values("date")
    )
    daily.to_csv(output, index=False, date_format="%Y-%m-%d", float_format="%.6f")
    metadata = {
        "sources": [str(historical_zip), str(current_xlsx)],
        "source_members": source_names,
        "filters": [
            "contract == T3PA (general EU allowances)",
            "status == successful",
            "positive auction volume and numeric auction price",
            "EAA3 aviation allowances excluded",
            "volume-weighted mean when multiple qualifying auctions share a date",
        ],
        "source_rows_read": source_rows,
        "output_rows": int(len(daily)),
        "first_date": daily["date"].min().date().isoformat(),
        "last_date": daily["date"].max().date().isoformat(),
        "unit": "EUR per tonne CO2 equivalent",
        "price_definition": "EEX primary-market EUA auction clearing price",
    }
    metadata_path.write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    return daily


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--historical-zip", type=Path, required=True)
    parser.add_argument("--current-xlsx", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--metadata", type=Path, default=DEFAULT_METADATA)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    for path in [args.historical_zip, args.current_xlsx]:
        if not path.exists() or path.stat().st_size == 0:
            raise SystemExit(f"Input file is missing or empty: {path}")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.metadata.parent.mkdir(parents=True, exist_ok=True)
    daily = build(args.historical_zip, args.current_xlsx, args.output, args.metadata)
    print(f"Built {len(daily)} daily EUA observations: {daily['date'].min().date()} to {daily['date'].max().date()}")
    print(f"Output: {args.output}")
    print(f"Metadata: {args.metadata}")


if __name__ == "__main__":
    main()
